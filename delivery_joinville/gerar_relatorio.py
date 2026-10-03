#!/usr/bin/env python3
"""Gera planilha (.xlsx) e PDF a partir do CSV do buscar_places.py.

Uso: python3 gerar_relatorio.py <csv> <categoria>
Categorias: sushi, pizza, hamburguer
"""
import csv
import re
import sys
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

CATEGORIAS = {
    "sushi": dict(
        titulo="Sushi e delivery em Joinville", foco="Sushi, japonesa, oriental e poke",
        termos='"sushi", "comida japonesa", "temakeria" e "poke"',
        tipos={"japanese_restaurant", "sushi_restaurant", "hawaiian_restaurant", "asian_restaurant",
               "japanese_izakaya_restaurant", "korean_restaurant", "chinese_restaurant"},
        nome=r"sushi|japan|japon|temak|poke|oriental|wok|nikkei|maki|lámen|lamen|yaki|asiát|hashi|kaze|koni|gokei|gendai|feng|xangai|chin"),
    "pizza": dict(
        titulo="Pizzarias e delivery em Joinville", foco="Pizzarias",
        termos='"pizzaria" e "pizza"',
        tipos={"pizza_restaurant", "pizza_delivery", "italian_restaurant"},
        nome=r"pizz|forno|forner|trattor|napol|ital"),
    "hamburguer": dict(
        titulo="Hamburguerias e delivery em Joinville", foco="Hamburguerias e lanches",
        termos='"hamburgueria", "hamburguer" e "lanchonete"',
        tipos={"hamburger_restaurant", "fast_food_restaurant", "sandwich_shop", "snack_bar", "american_restaurant", "hot_dog_restaurant"},
        nome=r"burg|lanch|smash|x-|dog|grill|sandu|food truck"),
}

PLATAFORMAS = [
    ("cardapioweb", "Cardápio Web"), ("goomer", "Goomer"), ("anota.ai", "Anota.ai"), ("pedir.delivery", "pedir.delivery"),
    ("ola.click", "Ola.click"), ("instadelivery", "InstaDelivery"), ("menudino", "MenuDino"), ("saipos", "Saipos"),
    ("ifood", "iFood"), ("lojalocal", "Loja Local"), ("comanda10", "Comanda10"), ("whatsmenu", "WhatsMenu"),
    ("mogo", "Mogo"), ("vocepede", "VocêPede"), ("cardapio.ai", "Cardapio.ai"), ("menuintegrado", "Menu Integrado"),
    ("ecta", "Ecta"), ("ireserve", "iReserve (reservas)"), ("neemo", "Neemo"), ("instagram", "Instagram"),
    ("facebook", "Facebook"), ("fb.me", "Facebook"), ("linktr.ee", "Linktree"), ("bio.site", "Bio.site"),
    ("wa.me", "WhatsApp"), ("whatsapp", "WhatsApp"),
]
SITUACAO = {"OPERATIONAL": "Aberto", "CLOSED_TEMPORARILY": "Fechado temporariamente", "CLOSED_PERMANENTLY": "Fechado"}
COLUNAS = ["Nome", "Bairro", "Endereço", "Site / link de pedido", "Plataforma", "Telefone", "Faz delivery?",
           "Nota Google", "Nº avaliações", "Situação", "Tipo (Google)", "Google Maps"]
LARGURAS = [38, 18, 55, 55, 18, 16, 13, 11, 13, 14, 22, 40]
DATA = "03/10/2026"


def plataforma(url):
    if not url:
        return "Sem site no Google"
    for chave, nome in PLATAFORMAS:
        if chave in url:
            return nome
    return "Site próprio"


def bairro(endereco):
    m = re.search(r" - ([^,]+), Joinville", endereco)
    return m.group(1) if m else ""


def linhas(registros):
    for r in registros:
        yield [r["nome"], bairro(r["endereco"]), r["endereco"], r["site"], plataforma(r["site"]), r["telefone"],
               r["delivery"] or "sem info", float(r["nota"]) if r["nota"] else None,
               int(r["avaliacoes"]) if r["avaliacoes"] else None, SITUACAO.get(r["situacao"], r["situacao"]),
               r["tipo"], r["google_maps"]]


def salvar_xlsx(caminho, cat, foco, outros):
    wb = Workbook()
    fonte, link = Font(name="Arial", size=10), Font(name="Arial", size=10, color="0563C1", underline="single")
    for ws, dados in ((wb.active, foco), (wb.create_sheet(), outros)):
        ws.append(COLUNAS)
        for c in ws[1]:
            c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="C0392B")
            c.alignment = Alignment(vertical="center")
        for linha in dados:
            ws.append(linha)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = fonte
            for i in (3, 11):
                if row[i].value:
                    row[i].hyperlink, row[i].font = row[i].value, link
        for i, w in enumerate(LARGURAS):
            ws.column_dimensions[chr(65 + i)].width = w
        ws.freeze_panes, ws.auto_filter.ref = "B2", ws.dimensions
    wb.active.title, wb.worksheets[1].title = cat["foco"][:31], "Outros resultados"
    notas = wb.create_sheet("Notas")
    for texto in [f"Fonte: Google Places API (New), busca de texto em {DATA}.",
                  f"Termos buscados: {cat['termos']} + Joinville, em 16 regiões da cidade.",
                  f'Aba "{cat["foco"]}": resultados cujo tipo no Google ou nome indica essa categoria.',
                  'Aba "Outros resultados": apareceram na busca, mas o cadastro no Google indica outra categoria.',
                  '"Faz delivery?" e "Site" vêm do cadastro no Google Maps; o restaurante pode ter outros canais não listados.',
                  "Plataforma: classificada automaticamente pelo endereço do site."]:
        notas.append([texto])
        notas.cell(notas.max_row, 1).font = fonte
    notas.column_dimensions["A"].width = 110
    wb.save(caminho)


def salvar_pdf(caminho, cat, foco, outros):
    pdfmetrics.registerFont(TTFont("DV", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"))
    pdfmetrics.registerFont(TTFont("DVB", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"))
    cel = ParagraphStyle("c", fontName="DV", fontSize=7, leading=8.5)
    lnk = ParagraphStyle("l", parent=cel, textColor=colors.HexColor("#0563C1"))
    cab = ParagraphStyle("h", fontName="DVB", fontSize=7.5, leading=9, textColor=colors.white)
    h1 = ParagraphStyle("h1", fontName="DVB", fontSize=16, leading=20, spaceAfter=4)
    h2 = ParagraphStyle("h2", fontName="DVB", fontSize=12, leading=15, spaceAfter=6, spaceBefore=4)
    txt = ParagraphStyle("t", fontName="DV", fontSize=8.5, leading=11)
    vermelho = colors.HexColor("#C0392B")

    def tabela(dados):
        t = [[Paragraph(h, cab) for h in ["#", "Nome", "Bairro", "Site / link de pedido", "Plataforma",
                                           "Telefone", "Entrega", "Nota", "Google Maps"]]]
        for i, (nome, bai, _, site, plat, tel, deliv, nota, aval, sit, _, gm) in enumerate(dados, 1):
            nome = escape(nome) + ("" if sit == "Aberto" else f' <font color="#C0392B">({sit})</font>')
            t.append([
                Paragraph(str(i), cel), Paragraph(nome, cel), Paragraph(escape(bai), cel),
                Paragraph(f'<link href="{escape(site)}">{escape(site.split("?")[0][:55])}</link>', lnk) if site else Paragraph("—", cel),
                Paragraph(escape(plat), cel), Paragraph(escape(tel), cel), Paragraph(escape(deliv), cel),
                Paragraph(f"{nota:.1f} ({aval})" if nota else "", cel),
                Paragraph(f'<link href="{escape(gm)}">abrir no mapa</link>', lnk) if gm else Paragraph("", cel),
            ])
        tab = Table(t, colWidths=[22, 150, 80, 200, 72, 72, 46, 52, 62], repeatRows=1)
        tab.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), vermelho), ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F6F6F6")]),
            ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#DDDDDD")),
            ("TOPPADDING", (0, 0), (-1, -1), 2.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5)]))
        return tab

    def rodape(c, doc):
        c.setFont("DV", 7)
        c.setFillColor(colors.grey)
        c.drawString(28, 16, f"{cat['titulo']} · Fonte: Google Places API, {DATA}")
        c.drawRightString(landscape(A4)[0] - 28, 16, f"Página {doc.page}")

    doc = SimpleDocTemplate(caminho, pagesize=landscape(A4), leftMargin=28, rightMargin=28,
                            topMargin=28, bottomMargin=30, title=cat["titulo"])
    story = [
        Paragraph(cat["titulo"], h1),
        Paragraph(f"Fonte: Google Places API, busca de {DATA} por {cat['termos']} em 16 regiões da cidade. "
                  "Site, telefone e delivery vêm do cadastro de cada restaurante no Google Maps; um restaurante "
                  "pode ter outros canais de pedido não cadastrados lá. Os links são clicáveis.", txt),
        Spacer(1, 8), Paragraph(f"{cat['foco']} ({len(foco)})", h2), tabela(foco),
    ]
    if outros:
        story += [PageBreak(), Paragraph(f"Outros resultados da busca ({len(outros)})", h2),
                  Paragraph("Apareceram na busca, mas o cadastro no Google indica outra categoria.", txt),
                  Spacer(1, 6), tabela(outros)]
    doc.build(story, onFirstPage=rodape, onLaterPages=rodape)


def main():
    csv_path, chave = sys.argv[1], sys.argv[2]
    cat = CATEGORIAS[chave]
    registros = list(csv.DictReader(open(csv_path, encoding="utf-8-sig")))
    padrao = re.compile(cat["nome"], re.I)
    eh_foco = [r["tipo"] in cat["tipos"] or bool(padrao.search(r["nome"])) for r in registros]
    foco = list(linhas([r for r, f in zip(registros, eh_foco) if f]))
    outros = list(linhas([r for r, f in zip(registros, eh_foco) if not f]))
    base = {"sushi": "sushi", "pizza": "pizzarias", "hamburguer": "hamburguerias"}[chave]
    salvar_xlsx(f"{base}_joinville_google_places.xlsx", cat, foco, outros)
    salvar_pdf(f"{base}_joinville_google_places.pdf", cat, foco, outros)
    print(f"{chave}: {len(foco)} na categoria, {len(outros)} outros")


if __name__ == "__main__":
    main()
