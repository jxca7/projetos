#!/usr/bin/env python3
"""Completa a lista de lojas Goomer / Cardápio Web com dados do Google Places e gera PDF + planilha.

Uso: GOOGLE_PLACES_API_KEY=... python3 relatorio_plataformas.py
"""
import csv
import difflib
import json
import os
import re
import unicodedata
import urllib.request
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

DATA = "03/10/2026"
CAMPOS = "places.displayName,places.formattedAddress,places.nationalPhoneNumber,places.rating,places.userRatingCount,places.googleMapsUri"


def normalizar(s):
    s = unicodedata.normalize("NFKD", s.lower()).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9 ]", " ", s)


def parecido(a, b):
    a, b = normalizar(a), normalizar(b)
    palavras = set(a.split()) - {"joinville", "de", "do", "da", "e", "the", "delivery"}
    return difflib.SequenceMatcher(None, a, b).ratio() > 0.6 or (palavras and len(palavras & set(b.split())) >= max(1, len(palavras) // 2))


def buscar(chave, nome):
    corpo = {"textQuery": f"{nome} Joinville SC", "pageSize": 3, "languageCode": "pt-BR"}
    req = urllib.request.Request("https://places.googleapis.com/v1/places:searchText", data=json.dumps(corpo).encode(),
                                 headers={"Content-Type": "application/json", "X-Goog-Api-Key": chave, "X-Goog-FieldMask": CAMPOS})
    for p in json.load(urllib.request.urlopen(req)).get("places", []):
        if "Joinville" in p.get("formattedAddress", "") and parecido(nome, p["displayName"]["text"]):
            return p
    return {}


def main():
    chave = os.environ["GOOGLE_PLACES_API_KEY"]
    lojas = list(csv.DictReader(open("plataformas_lista.csv", encoding="utf-8")))
    for l in lojas:
        p = buscar(chave, l["nome"])
        end = p.get("formattedAddress", "").replace(", Brasil", "")
        m = re.search(r"([^,]+), Joinville", end)
        l.update(endereco=end, bairro=m.group(1).split(" - ")[-1].strip() if m else "", telefone=p.get("nationalPhoneNumber", ""),
                 nota=f'{p["rating"]:.1f} ({p.get("userRatingCount", 0)})' if p.get("rating") else "",
                 maps=p.get("googleMapsUri", ""))

    # Planilha
    wb = Workbook()
    ws = wb.active
    ws.title = "Goomer e Cardápio Web"
    cab = ["Plataforma", "Categoria", "Nome", "Link do cardápio", "Bairro", "Endereço", "Telefone", "Nota Google", "Google Maps", "Observação"]
    ws.append(cab)
    for c in ws[1]:
        c.font, c.fill = Font(name="Arial", bold=True, color="FFFFFF"), PatternFill("solid", fgColor="C0392B")
    for l in lojas:
        ws.append([l["plataforma"], l["categoria"], l["nome"], l["link"], l["bairro"], l["endereco"], l["telefone"], l["nota"], l["maps"], l["obs"]])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name="Arial", size=10)
        for i in (3, 8):
            if row[i].value:
                row[i].hyperlink, row[i].font = row[i].value, Font(name="Arial", size=10, color="0563C1", underline="single")
    for col, w in zip("ABCDEFGHIJ", [14, 24, 34, 48, 18, 50, 16, 12, 40, 45]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes, ws.auto_filter.ref = "D2", ws.dimensions
    wb.save("goomer_cardapioweb_joinville.xlsx")

    # PDF
    pdfmetrics.registerFont(TTFont("DV", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"))
    pdfmetrics.registerFont(TTFont("DVB", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"))
    cel = ParagraphStyle("c", fontName="DV", fontSize=7.5, leading=9)
    lnk = ParagraphStyle("l", parent=cel, textColor=colors.HexColor("#0563C1"))
    hdr = ParagraphStyle("h", fontName="DVB", fontSize=8, leading=10, textColor=colors.white)
    h1 = ParagraphStyle("h1", fontName="DVB", fontSize=16, leading=20, spaceAfter=4)
    h2 = ParagraphStyle("h2", fontName="DVB", fontSize=12.5, leading=16, spaceBefore=10, spaceAfter=4, keepWithNext=1)
    h3 = ParagraphStyle("h3", fontName="DVB", fontSize=9.5, leading=12, spaceBefore=6, spaceAfter=3, keepWithNext=1)
    txt = ParagraphStyle("t", fontName="DV", fontSize=8.5, leading=11)

    def tabela(itens):
        dados = [[Paragraph(h, hdr) for h in ["#", "Nome", "Link do cardápio", "Bairro", "Telefone", "Nota", "Observação"]]]
        for i, l in enumerate(itens, 1):
            dados.append([Paragraph(str(i), cel), Paragraph(escape(l["nome"]), cel),
                          Paragraph(f'<link href="{escape(l["link"])}">{escape(l["link"])}</link>', lnk),
                          Paragraph(escape(l["bairro"]), cel), Paragraph(escape(l["telefone"]), cel),
                          Paragraph(escape(l["nota"]), cel), Paragraph(escape(l["obs"]), cel)])
        t = Table(dados, colWidths=[25, 145, 215, 85, 82, 55, 178], repeatRows=1)
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#C0392B")), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                               ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F6F6F6")]),
                               ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#DDDDDD")),
                               ("TOPPADDING", (0, 0), (-1, -1), 2.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5)]))
        return t

    def rodape(c, doc):
        c.setFont("DV", 7)
        c.setFillColor(colors.grey)
        c.drawString(28, 16, f"Goomer e Cardápio Web em Joinville · {DATA}")
        c.drawRightString(landscape(A4)[0] - 28, 16, f"Página {doc.page}")

    story = [Paragraph("Restaurantes de Joinville no Goomer e no Cardápio Web", h1),
             Paragraph(f"Levantamento de {DATA}: links encontrados no Google Places (busca de sushi, pizzaria e hamburgueria) "
                       "e em buscas na web pelas páginas \"Cardápio e Delivery em Joinville\" (Goomer) e \"Plataforma fornecida por "
                       "Cardápio Web\". Bairro, telefone e nota vêm do Google Maps; ficam em branco quando o restaurante não foi "
                       "encontrado lá com segurança. Os links são clicáveis.", txt)]
    for plat in ("Cardápio Web", "Goomer"):
        grupo = [l for l in lojas if l["plataforma"] == plat]
        story.append(Paragraph(f"{plat} ({len(grupo)})", h2))
        for cat in dict.fromkeys(l["categoria"] for l in grupo):
            itens = [l for l in grupo if l["categoria"] == cat]
            story += [Paragraph(f"{cat} ({len(itens)})", h3), tabela(itens)]
    doc = SimpleDocTemplate("goomer_cardapioweb_joinville.pdf", pagesize=landscape(A4), leftMargin=28, rightMargin=28,
                            topMargin=28, bottomMargin=30, title="Goomer e Cardápio Web em Joinville")
    doc.build(story, onFirstPage=rodape, onLaterPages=rodape)
    print(f"{len(lojas)} lojas; com dados do Google: {sum(1 for l in lojas if l['endereco'])}")


if __name__ == "__main__":
    main()
